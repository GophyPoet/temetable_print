"""
Excel parser module.

Reads .xlsx files, resolves cross-sheet formula references,
extracts schedule data with merged cell information, and filters sheets.
"""

import re
from dataclasses import dataclass, field
from typing import Optional

import openpyxl
from openpyxl.utils import get_column_letter

from . import config


@dataclass
class CellData:
    """Represents a single cell's content and formatting."""
    value: Optional[str] = None
    is_merged: bool = False
    merge_width: int = 1   # colspan
    merge_height: int = 1  # rowspan
    is_merged_slave: bool = False  # True if this cell is hidden by a merge
    bold: bool = False
    alignment: str = "center"  # left, center, right
    vertical_alignment: str = "middle"  # top, middle, bottom


@dataclass
class SheetData:
    """Parsed data from one worksheet."""
    name: str
    title: str = ""
    rows: list = field(default_factory=list)  # list of list of CellData
    num_rows: int = 0
    num_cols: int = 0
    source_file: str = ""
    teacher_row: Optional[list] = None
    skipped: bool = False
    skip_reason: str = ""


def should_skip_sheet(sheet_name: str) -> tuple[bool, str]:
    """Check if a sheet should be skipped based on config patterns."""
    name_lower = sheet_name.lower().strip()

    for pattern in config.SKIP_SHEET_PATTERNS:
        if pattern.lower() in name_lower:
            return True, f"Matches skip pattern '{pattern}'"

    if config.INCLUDE_SHEET_PATTERNS:
        for pattern in config.INCLUDE_SHEET_PATTERNS:
            if pattern.lower() in name_lower:
                return False, ""
        return True, "Does not match any include pattern"

    return False, ""


def resolve_formula(formula: str, workbook) -> Optional[str]:
    """
    Resolve simple cross-sheet cell references like =Звонки!B2.

    Only handles direct cell references (not complex formulas).
    """
    if not formula or not isinstance(formula, str):
        return None

    formula = formula.strip()
    if not formula.startswith("="):
        return None

    # Match pattern: =SheetName!CellRef
    match = re.match(r"^=([^!]+)!([A-Z]+\d+)$", formula[0:].strip(), re.IGNORECASE)
    if not match:
        # Try with quotes around sheet name: ='Sheet Name'!A1
        match = re.match(r"^='([^']+)'!([A-Z]+\d+)$", formula.strip(), re.IGNORECASE)

    if not match:
        return None

    sheet_name = match.group(1)
    cell_ref = match.group(2)

    try:
        if sheet_name in workbook.sheetnames:
            ws = workbook[sheet_name]
            cell = ws[cell_ref]
            if cell.value is not None:
                return str(cell.value)
    except Exception:
        pass

    return None


def get_merged_cell_info(ws) -> dict:
    """
    Build a map of merged cells.

    Returns dict: (row, col) -> {
        'master': bool,
        'colspan': int,
        'rowspan': int,
        'master_row': int,
        'master_col': int
    }
    """
    merge_map = {}

    for merged_range in ws.merged_cells.ranges:
        min_row = merged_range.min_row
        max_row = merged_range.max_row
        min_col = merged_range.min_col
        max_col = merged_range.max_col

        colspan = max_col - min_col + 1
        rowspan = max_row - min_row + 1

        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if row == min_row and col == min_col:
                    merge_map[(row, col)] = {
                        "master": True,
                        "colspan": colspan,
                        "rowspan": rowspan,
                    }
                else:
                    merge_map[(row, col)] = {
                        "master": False,
                        "colspan": 1,
                        "rowspan": 1,
                        "master_row": min_row,
                        "master_col": min_col,
                    }

    return merge_map


def parse_cell_value(cell, workbook) -> str:
    """Extract cell value, resolving formulas if needed."""
    value = cell.value

    if value is None:
        return ""

    if isinstance(value, str) and value.startswith("="):
        resolved = resolve_formula(value, workbook)
        if resolved is not None:
            return resolved
        return ""  # Unresolvable formula — show empty

    return str(value).strip()


def get_cell_alignment(cell) -> tuple[str, str]:
    """Extract alignment from cell formatting."""
    h_align = "center"
    v_align = "middle"

    if cell.alignment:
        if cell.alignment.horizontal:
            h_align = cell.alignment.horizontal
        if cell.alignment.vertical:
            v_align = cell.alignment.vertical
            if v_align == "center":
                v_align = "middle"

    return h_align, v_align


def parse_sheet(ws, workbook, source_file: str) -> SheetData:
    """Parse a single worksheet into SheetData."""
    sheet = SheetData(
        name=ws.title,
        source_file=source_file,
        num_rows=ws.max_row,
        num_cols=ws.max_column,
    )

    skip, reason = should_skip_sheet(ws.title)
    if skip:
        sheet.skipped = True
        sheet.skip_reason = reason
        return sheet

    if ws.max_row is None or ws.max_row < 3:
        sheet.skipped = True
        sheet.skip_reason = "Sheet has fewer than 3 rows — not a valid schedule"
        return sheet

    merge_map = get_merged_cell_info(ws)

    rows = []
    for row_idx in range(1, ws.max_row + 1):
        row_cells = []
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            merge_info = merge_map.get((row_idx, col_idx))

            if merge_info and not merge_info["master"]:
                # This cell is hidden by a merge — skip in rendering
                cd = CellData(
                    value="",
                    is_merged_slave=True,
                )
                row_cells.append(cd)
                continue

            value = parse_cell_value(cell, workbook)
            h_align, v_align = get_cell_alignment(cell)

            is_bold = False
            if cell.font and cell.font.bold:
                is_bold = True

            cd = CellData(
                value=value,
                is_merged=bool(merge_info and merge_info["master"]),
                merge_width=merge_info["colspan"] if merge_info and merge_info["master"] else 1,
                merge_height=merge_info["rowspan"] if merge_info and merge_info["master"] else 1,
                is_merged_slave=False,
                bold=is_bold,
                alignment=h_align,
                vertical_alignment=v_align,
            )
            row_cells.append(cd)

        rows.append(row_cells)

    sheet.rows = rows

    # Extract title from first row (usually merged across all columns)
    if rows and rows[0]:
        for cell in rows[0]:
            if cell.value and not cell.is_merged_slave:
                sheet.title = cell.value
                break

    return sheet


def parse_workbook(filepath: str, filename: str) -> list[SheetData]:
    """
    Parse an Excel workbook and return list of SheetData.

    Each sheet is parsed independently. Helper sheets are marked as skipped.
    """
    try:
        wb = openpyxl.load_workbook(filepath, data_only=False)
    except Exception as e:
        return [SheetData(
            name=filename,
            source_file=filename,
            skipped=True,
            skip_reason=f"Failed to open file: {str(e)}",
        )]

    sheets = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_data = parse_sheet(ws, wb, filename)
        sheets.append(sheet_data)

    wb.close()
    return sheets
